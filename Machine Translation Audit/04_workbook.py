"""Step 4: build one rater workbook per rater from sample/sample_items.jsonl.

Usage: python 04_workbook.py RaterA RaterB [RaterC ...]
Each workbook has a Rubric sheet and an Items sheet (item order shuffled per rater, seed derived from the name).
Rating columns use drop-down validation: rating in {OK, MINOR, MAJOR}; task_valid in {YES, NO, UNSURE}.
Raters never see idx or the other raters' sheets.
"""
import os, sys, json, random
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from audit_config import SAMPLE_DIR, SEED

RUBRIC = [
 ("What you are rating", "Whether the Turkish text faithfully conveys the English source, and whether the task's gold information still holds for the Turkish version."),
 ("OK", "Meaning preserved; fluent, idiomatic Turkish. Terminology appropriate for the domain."),
 ("MINOR", "Meaning preserved but the translation is flawed: over-literal or awkward phrasing, wrong register, inconsistent terminology, an English term left untranslated where a Turkish one exists, small stylistic additions that do not change content."),
 ("MAJOR", "Meaning changed: content added or omitted (including units, numbers, entities, qualifiers), negation or polarity flipped, a technical term mistranslated so that the content differs, a question that now asks something else."),
 ("task_valid", "YES if the gold information in 'task_context' still holds for the Turkish version (the gold label is still correct for the Turkish sentence(s); the Turkish question still identifies the gold code/answer). NO if the translation broke it. UNSURE if you cannot judge."),
 ("Code", "In code-related items, code is untouched by design; rate only the natural-language part. Identifiers, library names and code terms left in English are NOT errors."),
 ("Localisation", "The translation prompt asked the model to 'localise'. Reformatting numbers (1.7 -> 1,7) is fine. Adding information absent from the source (e.g. units, explanations) is MAJOR if it changes content, MINOR if purely cosmetic."),
 ("Independence", "Please do not discuss items with the other rater(s) before submitting. Use the comment column freely; it is read."),
]
COLS = [("item_id", 12), ("dataset", 14), ("rated_field", 16), ("english_source", 60), ("turkish_translation", 60),
        ("task_context", 50), ("rating", 10), ("task_valid", 11), ("comment", 40)]

def build(rater, items):
    wb = Workbook(); ws = wb.active; ws.title = "Rubric"
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 120
    ws.append(["Field", "Instruction"]); ws["A1"].font = ws["B1"].font = Font(bold=True)
    for k, v in RUBRIC:
        ws.append([k, v]); ws.cell(ws.max_row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    wi = wb.create_sheet("Items")
    for j, (h, w) in enumerate(COLS, 1):
        c = wi.cell(1, j, h); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDDDDD")
        wi.column_dimensions[get_column_letter(j)].width = w
    order = items[:]; random.Random(f"{SEED}-{rater}").shuffle(order)
    for i, it in enumerate(order, 2):
        vals = [it["item_id"], it["dataset"], it["rated_field"], it["english_source"], it["turkish_translation"], it["task_context"], "", "", ""]
        for j, v in enumerate(vals, 1):
            c = wi.cell(i, j, v); c.alignment = Alignment(wrap_text=True, vertical="top")
        for j in (7, 8):
            wi.cell(i, j).fill = PatternFill("solid", fgColor="FFF2CC")
    n = len(order) + 1
    dv1 = DataValidation(type="list", formula1='"OK,MINOR,MAJOR"', allow_blank=True); dv1.add(f"G2:G{n}")
    dv2 = DataValidation(type="list", formula1='"YES,NO,UNSURE"', allow_blank=True); dv2.add(f"H2:H{n}")
    wi.add_data_validation(dv1); wi.add_data_validation(dv2)
    wi.freeze_panes = "D2"
    path = os.path.join(SAMPLE_DIR, f"rater_{rater}.xlsx"); wb.save(path); return path

if __name__ == "__main__":
    raters = sys.argv[1:] or ["A", "B"]
    items = [json.loads(l) for l in open(os.path.join(SAMPLE_DIR, "sample_items.jsonl"), encoding="utf-8")]
    for r in raters:
        print("wrote", build(r, items), f"({len(items)} items)")
